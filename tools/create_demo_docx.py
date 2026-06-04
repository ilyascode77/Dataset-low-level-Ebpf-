from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/demo-encadrante-pfe.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(9)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True)
        set_cell_shading(hdr[i], "EAF2F8")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(35, 35, 35)
    return p


def add_screenshot_box(doc, label):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8F9FA")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(label)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(90, 90, 90)
    for _ in range(3):
        cell.add_paragraph("")
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [
        ("Title", 20, "1F4E79"),
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 11, "333333"),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "PFE - Dataset low-level eBPF / Kubernetes / Tetragon"
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Document de démonstration - Encadrante"
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
setup_styles(doc)
add_header_footer(doc)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Démonstration PFE - Dataset Low-Level eBPF pour Détection Ransomware-like dans Kubernetes")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Phase 1 validée : lab Kubernetes k3s, application bénigne, volumes PV/PVC, Tetragon/eBPF et génération de captures low-level")
run.font.name = "Arial"
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(80, 80, 80)

add_table(
    doc,
    ["Élément", "Valeur"],
    [
        ["Projet", "PFE cybersécurité cloud-native"],
        ["Environnement", "VM Ubuntu + k3s single-node"],
        ["Capteur principal", "Tetragon/eBPF"],
        ["Application bénigne", "Online Boutique - Google microservices-demo"],
        ["Objectif", "Générer un dataset low-level benign/malicious pour la détection ransomware-like"],
    ],
    widths=[2.2, 4.8],
)

doc.add_heading("1. Objectif Du Projet", level=1)
doc.add_paragraph(
    "L’objectif du projet est de construire un environnement cloud-native contrôlé basé sur Kubernetes afin de générer un dataset low-level pour la détection de comportements ransomware-like."
)
add_bullets(
    doc,
    [
        "Kubernetes k3s comme plateforme cloud-native locale.",
        "Online Boutique comme application e-commerce bénigne.",
        "Volumes persistants hostPath comme surfaces de données observables.",
        "Tetragon/eBPF pour collecter les événements low-level.",
        "Workloads I/O bénins pour produire une baseline normale.",
        "Scénarios attack-lab contrôlés pour simuler un comportement ransomware-like.",
    ],
)
doc.add_paragraph("L’objectif final est de produire un fichier dataset_lowlevel.csv contenant deux classes : benign et malicious.")

doc.add_heading("2. Architecture Générale", level=1)
add_code(
    doc,
    """VM1 - Ubuntu Server / Kubernetes k3s
|
+-- Namespace benign-lab
|   +-- Online Boutique Google microservices-demo
|   +-- fio workload
|   +-- file/data seeding workload
|   +-- PVC benign
|
+-- Namespace attack-lab
|   +-- suspicious-toolbox
|   +-- ransomware-like simulation
|   +-- PVC attack
|
+-- Namespace security
|   +-- Tetragon eBPF
|
+-- hostPath storage
    +-- /srv/pfe-lab/ecommerce
    +-- /srv/pfe-lab/fio
    +-- /srv/pfe-lab/filebench
    +-- /srv/pfe-lab/shared"""
)
add_table(
    doc,
    ["Composant", "Rôle"],
    [
        ["k3s", "Cluster Kubernetes local à un nœud."],
        ["benign-lab", "Contient les comportements normaux."],
        ["attack-lab", "Contient les scénarios ransomware-like contrôlés."],
        ["Tetragon", "Capteur eBPF pour événements low-level."],
        ["PV/PVC hostPath", "Surfaces de fichiers observables."],
        ["Online Boutique", "Application e-commerce réaliste."],
        ["fio", "Génération d’activité disque bénigne."],
        ["suspicious-toolbox", "Pod utilisé pour simuler l’attaque."],
    ],
    widths=[2.1, 4.9],
)

doc.add_heading("3. Justification Technique", level=1)
doc.add_heading("3.1 Pourquoi k3s + hostPath ?", level=2)
doc.add_paragraph(
    "Le choix de k3s permet d’avoir un cluster Kubernetes léger, réaliste et adapté à une VM étudiante. Le choix de hostPath est volontaire car le projet se fait sur un cluster single-node."
)
add_bullets(
    doc,
    [
        "Créer des volumes persistants simples.",
        "Observer les fichiers directement sur la VM.",
        "Générer des traces eBPF reproductibles.",
        "Simuler des accès bénins et malveillants aux mêmes surfaces de données.",
    ],
)
doc.add_paragraph(
    "Phrase rapport : Dans ce prototype de recherche, les volumes hostPath sont utilisés pour créer des surfaces de fichiers contrôlées et observables dans un cluster k3s single-node. Ce choix permet de générer une télémétrie low-level reproductible via Tetragon/eBPF sans dépendre d’un stockage cloud externe."
)

doc.add_heading("4. Vérification Du Lab", level=1)
doc.add_heading("4.1 Cluster Kubernetes", level=2)
add_code(doc, "kubectl get nodes -o wide\nkubectl get pods -A")
add_screenshot_box(doc, "[SCREENSHOT 1 - kubectl get nodes -o wide]")
add_screenshot_box(doc, "[SCREENSHOT 2 - kubectl get pods -A]")

doc.add_heading("4.2 Namespaces", level=2)
add_code(doc, "kubectl get ns")
add_screenshot_box(doc, "[SCREENSHOT 3 - namespaces Kubernetes]")

doc.add_heading("5. Application Bénigne : Online Boutique", level=1)
doc.add_paragraph("Online Boutique est déployée dans le namespace benign-lab.")
add_code(doc, "kubectl get pods -n benign-lab\nkubectl get svc -n benign-lab")
doc.add_paragraph("Accès à l’application :")
add_code(doc, "kubectl port-forward -n benign-lab svc/frontend 8080:80")
add_code(doc, "http://127.0.0.1:8080")
add_screenshot_box(doc, "[SCREENSHOT 4 - pods Online Boutique Running]")
add_screenshot_box(doc, "[SCREENSHOT 5 - interface Online Boutique]")

doc.add_heading("6. Stockage PV/PVC", level=1)
doc.add_heading("6.1 Structure hostPath", level=2)
add_code(doc, "/srv/pfe-lab/ecommerce\n/srv/pfe-lab/fio\n/srv/pfe-lab/filebench\n/srv/pfe-lab/shared")
add_code(doc, "sudo find /srv/pfe-lab -maxdepth 2 -type d")
doc.add_heading("6.2 Vérification des PVC", level=2)
add_code(doc, "kubectl get pvc -n benign-lab\nkubectl get pvc -n attack-lab\nkubectl get pv")
add_screenshot_box(doc, "[SCREENSHOT 6 - PVC benign-lab Bound]")
add_screenshot_box(doc, "[SCREENSHOT 7 - PVC attack-lab Bound]")
add_table(
    doc,
    ["PV", "PVC", "Namespace", "Usage"],
    [
        ["pv-ecommerce-benign", "pvc-ecommerce", "benign-lab", "Données e-commerce"],
        ["pv-fio", "pvc-fio", "benign-lab", "Workload I/O"],
        ["pv-filebench", "pvc-filebench", "benign-lab", "Workload fichiers"],
        ["pv-shared-benign", "pvc-shared-benign", "benign-lab", "Volume partagé bénin"],
        ["pv-ecommerce-attack", "pvc-ecommerce-attack", "attack-lab", "Accès attack aux données e-commerce"],
        ["pv-shared-attack", "pvc-shared-attack", "attack-lab", "Accès attack au volume partagé"],
    ],
    widths=[1.8, 1.8, 1.3, 2.1],
)

doc.add_heading("7. Génération De Données Bénignes", level=1)
doc.add_heading("7.1 Seed e-commerce data", level=2)
add_code(doc, "kubectl apply -f k8s/benign-workloads/io/seed-ecommerce-data-job.yaml\nkubectl wait -n benign-lab --for=condition=complete job/seed-ecommerce-data --timeout=120s\nkubectl logs -n benign-lab job/seed-ecommerce-data")
doc.add_heading("7.2 fio workload", level=2)
add_code(doc, "kubectl apply -f k8s/benign-workloads/io/fio-configmap.yaml\nkubectl apply -f k8s/benign-workloads/io/fio-job.yaml\nkubectl wait -n benign-lab --for=condition=complete job/fio-workload --timeout=300s\nkubectl logs -n benign-lab job/fio-workload")
add_code(doc, "kubectl get jobs -n benign-lab")
add_screenshot_box(doc, "[SCREENSHOT 8 - jobs benign Complete]")

doc.add_heading("8. Tetragon/eBPF", level=1)
doc.add_paragraph("Tetragon est installé dans le namespace security.")
add_code(doc, "kubectl get pods -n security\nkubectl get ds -n security\nkubectl get tracingpolicies")
add_screenshot_box(doc, "[SCREENSHOT 9 - Tetragon Running]")
add_screenshot_box(doc, "[SCREENSHOT 10 - TracingPolicies]")
add_table(
    doc,
    ["Policy", "Rôle"],
    [
        ["pfe-network-observability", "Capture tcp_connect, tcp_sendmsg, tcp_close."],
        ["pfe-sensitive-file-access", "Capture accès fichiers sensibles."],
        ["pfe-ransomware-like-file-activity", "Capture accès aux fichiers victimes."],
    ],
    widths=[2.5, 4.5],
)

doc.add_heading("9. Capture Des Logs Bénins", level=1)
add_code(doc, "mkdir -p dataset/raw dataset/generated")
add_code(doc, "mv benign_capture.json dataset/raw/benign_capture_lowlevel.jsonl")
add_code(doc, "timeout 60 kubectl exec -n security ds/tetragon -c tetragon -- tetra getevents -o json --namespace benign-lab > dataset/raw/benign_capture_lowlevel.jsonl")
add_code(doc, "for i in $(seq 1 30); do curl -s http://127.0.0.1:8080 >/dev/null || true; done")
add_code(doc, "wc -l dataset/raw/benign_capture_lowlevel.jsonl\nhead -n 2 dataset/raw/benign_capture_lowlevel.jsonl")
add_screenshot_box(doc, "[SCREENSHOT 11 - benign_capture_lowlevel.jsonl]")

doc.add_heading("10. Simulation Attack Ransomware-like", level=1)
add_code(doc, "kubectl apply -f k8s/attack-lab/suspicious-toolbox.yaml\nkubectl wait -n attack-lab --for=condition=Ready pod/suspicious-toolbox --timeout=120s\nkubectl get pods -n attack-lab")
doc.add_paragraph("Terminal 1 :")
add_code(doc, "timeout 60 kubectl exec -n security ds/tetragon -c tetragon -- tetra getevents -o json --namespace attack-lab > dataset/raw/attack_capture_lowlevel.jsonl")
doc.add_paragraph("Terminal 2 :")
add_code(doc, "kubectl exec -n attack-lab suspicious-toolbox -- /bin/sh -c 'rm -rf /victims && mkdir -p /victims && for i in $(seq 1 20); do echo test > /victims/file-$i.txt; cp /victims/file-$i.txt /victims/file-$i.locked; done'")
add_code(doc, "wc -l dataset/raw/attack_capture_lowlevel.jsonl\nhead -n 2 dataset/raw/attack_capture_lowlevel.jsonl")
add_screenshot_box(doc, "[SCREENSHOT 12 - attack_capture_lowlevel.jsonl]")
doc.add_paragraph(
    "Le scénario malveillant ne contient aucun malware réel. Il reproduit uniquement des comportements ransomware-like observables : création massive de fichiers, copie vers une extension .locked, accès répétés aux fichiers et exécution de processus suspects dans un namespace isolé."
)

doc.add_heading("11. Transformation En Dataset CSV", level=1)
doc.add_heading("11.1 Parser JSONL vers CSV", level=2)
add_code(doc, "nano scripts/tetragon-jsonl-to-csv.py\nchmod +x scripts/tetragon-jsonl-to-csv.py")
doc.add_heading("11.2 Convertir benign", level=2)
add_code(doc, "python3 scripts/tetragon-jsonl-to-csv.py \\\n  --input dataset/raw/benign_capture_lowlevel.jsonl \\\n  --output dataset/generated/benign_lowlevel.csv \\\n  --label benign \\\n  --scenario benign-cloud-native-baseline")
doc.add_heading("11.3 Convertir attack", level=2)
add_code(doc, "python3 scripts/tetragon-jsonl-to-csv.py \\\n  --input dataset/raw/attack_capture_lowlevel.jsonl \\\n  --output dataset/generated/attack_lowlevel.csv \\\n  --label malicious \\\n  --scenario ransomware-like-file-copy-lock")
doc.add_heading("11.4 Fusionner", level=2)
add_code(doc, "head -n 1 dataset/generated/benign_lowlevel.csv > dataset/generated/dataset_lowlevel.csv\ntail -n +2 dataset/generated/benign_lowlevel.csv >> dataset/generated/dataset_lowlevel.csv\ntail -n +2 dataset/generated/attack_lowlevel.csv >> dataset/generated/dataset_lowlevel.csv")
add_code(doc, "wc -l dataset/generated/*.csv\nhead -n 5 dataset/generated/dataset_lowlevel.csv")
add_screenshot_box(doc, "[SCREENSHOT 13 - CSV generated]")
add_screenshot_box(doc, "[SCREENSHOT 14 - dataset_lowlevel.csv head]")

doc.add_heading("12. Attributs Du Dataset Low-Level", level=1)
add_table(
    doc,
    ["Attribut", "Description"],
    [
        ["timestamp", "Date de l’événement."],
        ["event_type", "Type d’événement Tetragon."],
        ["function_name", "Fonction kernel observée."],
        ["policy_name", "Nom de la policy Tetragon."],
        ["node_name", "Nom du nœud Kubernetes."],
        ["namespace", "Namespace Kubernetes."],
        ["pod", "Nom du pod."],
        ["workload", "Workload Kubernetes."],
        ["container / image", "Contexte container."],
        ["binary / arguments", "Processus exécuté et arguments."],
        ["src_ip / dst_ip", "Informations réseau."],
        ["file_path", "Chemin fichier observé."],
        ["label", "benign ou malicious."],
        ["scenario", "Nom du scénario de capture."],
    ],
    widths=[2.4, 4.6],
)

doc.add_heading("13. Résultat Attendu", level=1)
doc.add_paragraph("À la fin de la démonstration, on obtient :")
add_code(doc, "dataset/generated/dataset_lowlevel.csv")
doc.add_paragraph("Ce fichier contient des lignes benign et malicious, issues d’événements low-level Tetragon/eBPF.")
add_screenshot_box(doc, "[SCREENSHOT 15 - dataset final]")

doc.add_heading("14. Prochaines Étapes", level=1)
add_bullets(
    doc,
    [
        "Ajouter un scénario ransomware-family1 plus complet utilisant pvc-ecommerce-attack et pvc-shared-attack.",
        "Ajouter Falco pour générer des alertes runtime.",
        "Envoyer les alertes vers Wazuh/SOC.",
        "Construire des règles de détection.",
        "Ajouter une VM Kali pour simuler un attaquant externe.",
        "Améliorer l’analyse du dataset avec Python/Jupyter.",
    ],
)

doc.add_heading("15. Conclusion", level=1)
doc.add_paragraph(
    "Cette première phase montre que l’environnement Kubernetes génère des événements low-level exploitables via eBPF/Tetragon. Les comportements bénins et ransomware-like sont capturés, labellisés et transformés en dataset CSV."
)
add_code(doc, "Cloud-native workload -> eBPF low-level telemetry -> dataset generation -> ransomware-like detection -> future SOC/XDR integration")

doc.save(OUT)
print(OUT)
